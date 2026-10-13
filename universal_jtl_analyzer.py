#!/usr/bin/env python3
"""Universal, rule-based JMeter CSV/XML JTL -> Excel analyzer (no AI).
Usage: python universal_jtl_analyzer.py input.jtl [-o report.xlsx] [--gap-seconds 120]
Requires: pip install pandas numpy openpyxl
"""
import argparse
import csv
import re
import xml.etree.ElementTree as ET
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ALIASES = {
    'timestamp': ['timestamp', 'time_stamp', 'starttime', 'start_time', 'ts'],
    'elapsed': ['elapsed', 'responsetime', 'response_time', 'duration', 'durationms', 't'],
    'label': ['label', 'samplerlabel', 'sampler_label', 'name', 'lb'],
    'success': ['success', 'successful', 'ok', 's'],
    'thread': ['threadname', 'thread_name', 'thread', 'tn'],
    'code': ['responsecode', 'response_code', 'statuscode', 'status_code', 'rc'],
    'message': ['responsemessage', 'response_message', 'rm'],
    'latency': ['latency', 'lt'],
    'connect': ['connect', 'connecttime', 'ct'],
    'bytes': ['bytes', 'by'],
    'url': ['url'],
    'sample_count': ['samplecount', 'sample_count', 'sc'],
    'error_count': ['errorcount', 'error_count', 'ec'],
    'all_threads': ['allthreads', 'all_threads', 'na'],
    'group_threads': ['grpthreads', 'groupthreads', 'ng'],
    'flow': ['flowname', 'flow_name', 'scenario', 'journey', 'businessflow'],
    'iteration': ['iterationid', 'iteration_id', 'iteration', 'loopid', 'loop_id'],
    'parent': ['isparent', 'parent', 'transactionparent'],
}

def norm(s):
    return re.sub(r'[^a-z0-9]', '', str(s).lower())

def read_input(path):
    with open(path, 'rb') as f:
        head = f.read(4096).lstrip()
    if head.startswith(b'<'):
        rows = []
        for _, el in ET.iterparse(path, events=('end',)):
            if el.tag in ('httpSample', 'sample'):
                row = dict(el.attrib)
                if 'url' not in row:
                    url = el.find('java.net.URL')
                    if url is not None and url.text:
                        row['url'] = url.text
                rows.append(row)
                el.clear()
        if not rows:
            raise ValueError('No JMeter XML sample/httpSample elements found')
        return pd.DataFrame(rows), 'XML'
    for enc in ('utf-8-sig', 'utf-16', 'latin-1'):
        try:
            with open(path, encoding=enc, newline='') as f:
                preview = f.read(8192)
            delimiter = csv.Sniffer().sniff(preview, delimiters=',;\t|').delimiter
            return pd.read_csv(path, sep=delimiter, encoding=enc, low_memory=False), 'CSV'
        except (UnicodeError, csv.Error, pd.errors.ParserError):
            continue
    raise ValueError('Unable to parse CSV JTL')

def normalize(df):
    lookup = {norm(c): c for c in df.columns}
    mapping = {}
    for key, candidates in ALIASES.items():
        for candidate in candidates:
            if norm(candidate) in lookup:
                mapping[key] = lookup[norm(candidate)]
                break
    for required in ('label', 'elapsed'):
        if required not in mapping:
            raise ValueError(f'Missing required field {required}. Available: {list(df.columns)}')
    out = pd.DataFrame(index=df.index)
    for key, col in mapping.items():
        out[key] = df[col]
    out['label'] = out['label'].fillna('(unlabeled)').astype(str)
    out['elapsed'] = pd.to_numeric(out['elapsed'], errors='coerce')
    out = out.dropna(subset=['elapsed']).copy()
    out = out[out['elapsed'] >= 0].copy()
    if out.empty:
        raise ValueError('No valid elapsed values')
    if 'timestamp' in out:
        raw = out['timestamp']
        numeric = pd.to_numeric(raw, errors='coerce')
        if numeric.notna().mean() >= 0.8:
            med = numeric.dropna().abs().median()
            # Epoch seconds vs epoch milliseconds; other numeric values treated as relative milliseconds.
            if 1e9 < med < 1e11:
                numeric = numeric * 1000
            out['timestamp'] = numeric
        else:
            parsed = pd.to_datetime(raw, errors='coerce', utc=True)
            out['timestamp'] = parsed.map(lambda v: v.timestamp() * 1000 if pd.notna(v) else np.nan)
    else:
        out['timestamp'] = np.nan
    for key, default in [('thread', 'unknown'), ('flow', 'Unspecified'), ('iteration', ''), ('code', ''), ('message', ''), ('url', '')]:
        if key not in out:
            out[key] = default
        out[key] = out[key].fillna(default).astype(str)
    if 'success' in out:
        raw = out['success'].astype(str).str.strip().str.lower()
        out['success'] = raw.map({'true': True, 'false': False, '1': True, '0': False, 'yes': True, 'no': False, 't': True, 'f': False})
    else:
        code = pd.to_numeric(out['code'], errors='coerce')
        out['success'] = code.between(200, 399).astype('boolean') if code.notna().any() else pd.Series(pd.NA, index=out.index, dtype='boolean')
    for key in ('sample_count', 'error_count', 'latency', 'connect', 'bytes', 'all_threads', 'group_threads'):
        if key not in out:
            out[key] = np.nan
        out[key] = pd.to_numeric(out[key], errors='coerce')
    out['sample_count'] = out['sample_count'].fillna(1).clip(lower=1)
    out['error_count'] = out['error_count'].where(out['error_count'].notna(), np.where(out['success'].eq(False).fillna(False), out['sample_count'], 0))
    out['end'] = out['timestamp'] + out['elapsed']
    out['_order'] = np.arange(len(out))
    out = out.sort_values(['timestamp', '_order'], na_position='last').reset_index(drop=True)
    return out, mapping

def pctl(s, q):
    return round(float(s.quantile(q)), 2) if len(s) else None

def stats(df, key):
    rows = []
    for name, g in df.groupby(key, dropna=False, sort=False):
        samples = int(g['sample_count'].sum())
        errors = int(g['error_count'].sum())
        times = g['elapsed']
        valid = g.dropna(subset=['timestamp', 'end'])
        span = (valid['end'].max() - valid['timestamp'].min()) / 1000 if len(valid) else np.nan
        rows.append({'Name': name, 'Recorded Rows': len(g), 'Samples': samples,
                     'Failures': errors, 'Error %': round(100 * errors / samples, 2) if samples else None,
                     'Avg ms': round(times.mean(), 2), 'Median ms': pctl(times, .5),
                     'Min ms': round(times.min(), 2), 'Max ms': round(times.max(), 2),
                     'P90 ms': pctl(times, .9), 'P95 ms': pctl(times, .95), 'P99 ms': pctl(times, .99),
                     'Observed TPS': round(samples / span, 3) if pd.notna(span) and span > 0 else None,
                     'P95 SLA': 'FAIL' if pctl(times, .95) > 3000 else 'PASS',
                     'Error SLA': 'FAIL' if errors / samples > .01 else 'PASS'})
    return pd.DataFrame(rows).sort_values('P95 ms', ascending=False)

def journeys(df, gap_seconds):
    """Explicit iteration IDs are authoritative; otherwise split on long inactivity only.
    A thread sequence is NOT asserted to be a complete business iteration.
    """
    if df['timestamp'].isna().all():
        return pd.DataFrame(), pd.DataFrame(), 'No timestamps: journey order cannot be verified.'
    events = []
    summary = []
    jid = 0
    explicit = df['iteration'].str.strip().ne('').any()
    for (flow, thread), g in df.groupby(['flow', 'thread'], sort=False):
        g = g.sort_values(['timestamp', '_order'])
        groups = []
        if explicit and g['iteration'].str.strip().ne('').all():
            groups = [(str(k), part) for k, part in g.groupby('iteration', sort=False)]
        else:
            # Group on idle gap only, not repeated labels: repeats can be valid within one journey.
            prior_end = g['end'].shift()
            gaps = (g['timestamp'] - prior_end).fillna(0)
            segment = (gaps > gap_seconds * 1000).cumsum()
            groups = [(str(k), part) for k, part in g.groupby(segment, sort=False)]
        for iteration, part in groups:
            jid += 1
            ordered = part.sort_values(['timestamp', '_order'])
            start, end = ordered['timestamp'].min(), ordered['end'].max()
            duration = end - start
            failures = int(ordered['error_count'].sum())
            summary.append({'Journey ID': jid, 'Flow': flow, 'Thread': thread, 'Iteration/Segment': iteration,
                            'Boundary Confidence': 'Explicit iteration' if explicit and ordered['iteration'].str.strip().ne('').all() else 'Heuristic thread segment',
                            'Steps': len(ordered), 'Unique Steps': ordered['label'].nunique(),
                            'Wall Duration ms': round(duration, 2), 'Sum Sample ms': round(ordered['elapsed'].sum(), 2),
                            'Failures': failures, 'Result': 'FAIL' if failures else 'PASS/UNKNOWN',
                            'Sequence': ' -> '.join(ordered['label'].astype(str).tolist())[:30000]})
            for step, (_, r) in enumerate(ordered.iterrows(), 1):
                events.append({'Journey ID': jid, 'Flow': flow, 'Thread': thread, 'Step': step,
                               'Label': r['label'], 'Start UTC': pd.to_datetime(r['timestamp'], unit='ms', utc=True).strftime('%Y-%m-%d %H:%M:%S.%f')[:-3] if pd.notna(r['timestamp']) else '',
                               'Elapsed ms': r['elapsed'], 'Success': 'UNKNOWN' if pd.isna(r['success']) else ('PASS' if r['success'] else 'FAIL'),
                               'Code': r['code'], 'URL': r['url']})
    note = 'Explicit iteration IDs used where present.' if explicit else 'Journey boundaries are heuristic; thread segments may contain multiple iterations or partial flows.'
    return pd.DataFrame(events), pd.DataFrame(summary), note

def interpretations(df, step, flow, journey, note):
    findings = []
    def add(level, area, evidence, meaning, action):
        findings.append({'Severity': level, 'Area': area, 'Evidence': evidence, 'Interpretation': meaning, 'Suggested Action': action})
    count = int(df['sample_count'].sum())
    err = int(df['error_count'].sum())
    if err:
        add('HIGH' if err / count > .01 else 'INFO', 'Reliability', f'{err}/{count} failed ({err/count:.2%})',
            'Some recorded samples failed.', 'Filter errors by label and response code; correlate with application logs.')
    else:
        add('INFO', 'Reliability', 'No recorded failures', 'No failures found in available success/error fields.', 'Validate functional assertions and JTL success recording.')
    for _, r in step.iterrows():
        if r['P95 ms'] > 3000:
            add('HIGH', 'Latency', f"{r['Name']}: P95 {r['P95 ms']} ms", 'P95 exceeds example 3000 ms SLA.', 'Check downstream dependencies, DB time, CPU, and load phase.')
        if r['Error %'] > 1:
            add('HIGH', 'Errors', f"{r['Name']}: {r['Error %']}% errors", 'Exceeds example 1% error threshold.', 'Inspect response codes and error messages.')
        if r['Max ms'] > max(1000, 3 * r['P95 ms']):
            add('MEDIUM', 'Outliers', f"{r['Name']}: max {r['Max ms']} ms vs P95 {r['P95 ms']} ms", 'Occasional latency spikes are present.', 'Inspect time series and backend events.')
    if len(journey):
        failures = (journey['Failures'] > 0).sum()
        add('INFO', 'Journeys', f'{len(journey)} reconstructed segments; {failures} have failed samples', note,
            'Add FlowName and IterationID to JTL for reliable business-flow boundaries.')
    else:
        add('WARNING', 'Journeys', 'Journey reconstruction unavailable', note, 'Enable timeStamp and threadName in JMeter JTL.')
    add('INFO', 'Data quality', 'Parent and child samples may both be recorded',
        'Parent transaction elapsed time can overlap child elapsed time; summed durations are not end-to-end time.',
        'Use wall-clock journey duration and exclude parent samples from step totals when identified.')
    if df['success'].isna().any():
        add('WARNING', 'Data quality', f"{df['success'].isna().sum()} rows have unknown success", 'Reliability estimates may be incomplete.', 'Save success and responseCode in JTL.')
    return pd.DataFrame(findings)

def write_excel(tables, path):
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        for name, data in tables.items():
            if data is None or data.empty:
                data = pd.DataFrame({'Information': ['No data available for this section']})
            data.to_excel(writer, sheet_name=name[:31], index=False)
    book = load_workbook(path)
    for sheet in book:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color='FFFFFF', bold=True)
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.alignment = Alignment(wrap_text=True)
        for col in sheet.columns:
            values = [len(str(c.value or '')) for c in list(col)[:2000]]
            sheet.column_dimensions[get_column_letter(col[0].column)].width = min(65, max(14, max(values, default=12) + 2))
    book.save(path)

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', help='JMeter .jtl, .csv or XML JTL')
    parser.add_argument('-o', '--output', default='JTL_Universal_Analysis.xlsx')
    parser.add_argument('--gap-seconds', type=float, default=120, help='Idle gap for heuristic journey segmentation')
    args = parser.parse_args()
    raw, fmt = read_input(args.input)
    df, mapping = normalize(raw)
    step = stats(df, 'label')
    flow = stats(df, 'flow')
    events, journey, note = journeys(df, args.gap_seconds)
    errors = df[df['error_count'] > 0].groupby(['label', 'code', 'message'], dropna=False).agg(Failed_Rows=('label', 'size'), Estimated_Failed_Samples=('error_count', 'sum')).reset_index()
    transitions = pd.DataFrame()
    if not events.empty:
        ordered = events.sort_values(['Journey ID', 'Step']).copy()
        ordered['Next Step'] = ordered.groupby('Journey ID')['Label'].shift(-1)
        transitions = ordered.dropna(subset=['Next Step']).groupby(['Flow', 'Label', 'Next Step']).size().reset_index(name='Observed Count').sort_values('Observed Count', ascending=False)
    total = int(df['sample_count'].sum())
    duration = (df['end'].max() - df['timestamp'].min()) / 1000 if df['timestamp'].notna().any() else np.nan
    summary = pd.DataFrame([
        ('Input file', str(Path(args.input).name)), ('Detected format', fmt), ('Recorded rows', len(df)),
        ('Estimated sample count', total), ('Distinct labels', df['label'].nunique()),
        ('Distinct threads', df['thread'].nunique()), ('Recorded failures', int(df['error_count'].sum())),
        ('Error %', round(100 * df['error_count'].sum() / total, 2)),
        ('Test wall duration seconds', round(duration, 2) if pd.notna(duration) else 'Unavailable'),
        ('Overall observed samples/sec', round(total / duration, 3) if pd.notna(duration) and duration > 0 else 'Unavailable'),
        ('Journey confidence', note),
        ('SLA assumptions', 'P95 <= 3000 ms; error rate <= 1%; configurable in script'),
        ('Interpretation caution', 'No AI; labels do not reveal unrecorded UI steps or true business semantics')
    ], columns=['Metric', 'Value'])
    quality = pd.DataFrame([{'Canonical Field': k, 'Source Column': mapping.get(k, 'NOT PRESENT')} for k in ALIASES])
    write_excel({'Executive Summary': summary, 'Findings': interpretations(df, step, flow, journey, note),
                 'Step Analysis': step, 'Flow Analysis': flow, 'Journey Summary': journey,
                 'Journey Steps': events, 'Step Transitions': transitions, 'Error Analysis': errors,
                 'Data Mapping': quality}, args.output)
    print(f'Analyzed {len(df):,} rows. Saved: {args.output}')
    print(note)

if __name__ == '__main__':
    main()
